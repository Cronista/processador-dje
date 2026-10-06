import streamlit as st
import pandas as pd
import fitz  # PyMuPDF
import os
import re
import numbers
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font

from analise_dje import (
    TERMOS_CARGOS, TERMOS_ACAO, TERMO_MATRICULA, CONTA, NAO_CONTA, VERIFICAR,
    listar_pdfs, ler_paginas_pdf, dividir_em_atos, analisar_documento, BuscaPortarias
)

# ==========================================
# CONFIGURAÇÃO DA PÁGINA E VARIÁVEIS GERAIS
# ==========================================
st.set_page_config(page_title="Processador de DJE", page_icon="📄", layout="wide")

# Pasta onde está o app.py: caminhos relativos e arquivos gerados ficam sempre aqui,
# independentemente de onde o Streamlit foi iniciado.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

MESES_EXTENSO = {
    1: "JANEIRO", 2: "FEVEREIRO", 3: "MARÇO", 4: "ABRIL",
    5: "MAIO", 6: "JUNHO", 7: "JULHO", 8: "AGOSTO",
    9: "SETEMBRO", 10: "OUTUBRO", 11: "NOVEMBRO", 12: "DEZEMBRO"
}

# ==========================================
# FUNÇÕES AUXILIARES
# ==========================================

def resolver_caminho(caminho):
    caminho = os.path.expanduser(caminho.strip())
    if os.path.isabs(caminho):
        return caminho
    return os.path.normpath(os.path.join(BASE_DIR, caminho))

def nome_com_timestamp(nome_base):
    nome, ext = os.path.splitext(nome_base)
    return f"{nome}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"

def listar_planilhas():
    """Planilhas .xlsx da pasta do sistema, da mais recente para a mais antiga."""
    arquivos = [f for f in os.listdir(BASE_DIR) if f.lower().endswith(".xlsx") and not f.startswith("~$")]
    return sorted(arquivos, key=lambda f: os.path.getmtime(os.path.join(BASE_DIR, f)), reverse=True)

def converter_data(valor):
    """Converte o valor da coluna DATA (texto dd/mm/aaaa ou data do Excel) em datetime. None se inválido."""
    if isinstance(valor, datetime):  # inclui pd.Timestamp
        return valor
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y")
    except ValueError:
        return None

def formatar_data(valor):
    data = converter_data(valor)
    return data.strftime("%d/%m/%Y") if data else str(valor)

def extrair_lista_paginas(valor):
    """Converte a coluna PÁGINAS em lista de números. Aceita '3, 5; 7', 12 ou 12.0 (célula editada no Excel)."""
    if pd.isna(valor):
        return []
    if isinstance(valor, numbers.Real):
        return [int(valor)]
    texto = str(valor).strip()
    if texto.upper() == "SEM":
        return []
    return [int(p) for p in re.split(r'[,;\s]+', texto) if p.isdigit()]

def ler_planilha_indice(caminho_planilha):
    if not caminho_planilha or not os.path.exists(caminho_planilha):
        return None, f"Planilha de índice não encontrada: {caminho_planilha or '(nenhuma selecionada)'}. Execute a Aba 1 primeiro."
    try:
        df = pd.read_excel(caminho_planilha)
    except Exception as e:
        return None, f"Não foi possível abrir a planilha {os.path.basename(caminho_planilha)}: {e}"
    faltando = [col for col in ("DATA", "PÁGINAS") if col not in df.columns]
    if faltando:
        return None, f"A planilha {os.path.basename(caminho_planilha)} não tem as colunas obrigatórias: {', '.join(faltando)}."
    return df, None

COLUNAS_TEXTO_LONGO = {"MOTIVO", "REFERÊNCIA", "TRECHO"}
CORES_RESULTADO = {CONTA: "C6EFCE", NAO_CONTA: "EDEDED", VERIFICAR: "FFEB9C"}

def estilizar_planilha(arquivo):
    wb = load_workbook(arquivo)
    cor_cabecalho = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fonte_cabecalho = Font(bold=True, name='Calibri')
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_texto = Alignment(horizontal="left", vertical="top", wrap_text=True)
    borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for ws in wb.worksheets:
        cabecalhos = [cell.value for cell in ws[1]]
        for cell in ws[1]:
            cell.fill = cor_cabecalho
            cell.font = fonte_cabecalho
            cell.alignment = alinhamento_centro
            cell.border = borda_fina

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell, coluna in zip(row, cabecalhos):
                cell.alignment = alinhamento_texto if coluna in COLUNAS_TEXTO_LONGO else alinhamento_centro
                cell.border = borda_fina
                if coluna == "RESULTADO":
                    # "VERIFICAR (dispensável)" usa a mesma cor de "VERIFICAR"
                    cor = CORES_RESULTADO.get(str(cell.value).split(" (")[0])
                    if cor:
                        cell.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column].width = min(max_length + 4, 60)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(arquivo)

# ==========================================
# FUNÇÕES DE PROCESSAMENTO
# ==========================================

def gerar_indice_excel(pasta_diarios, pasta_acervo, arquivo_saida):
    """Retorna (índice ou None, detalhes por ato, mensagem, lista de erros por arquivo).

    Uma página entra no índice quando tem portaria de cargo/função ou apostila (ou "tornar sem efeito")
    com efeito financeiro. Apostilas que não puderam ser confirmadas também entram, marcadas em
    "PÁGINAS A VERIFICAR" e explicadas na aba "Detalhes" da planilha.
    """
    if not os.path.exists(pasta_diarios):
        return None, None, f"A pasta '{pasta_diarios}' não existe.", []

    arquivos_pdf = listar_pdfs(pasta_diarios)
    if not arquivos_pdf:
        return None, None, "Nenhum PDF encontrado na pasta de íntegras.", []

    # As portarias citadas pelas apostilas são procuradas nos DJEs da própria pasta e no acervo
    busca = BuscaPortarias([pasta_diarios, pasta_acervo])
    resultados = []
    detalhes = []
    erros = []
    barra_progresso = st.progress(0)

    for i, arquivo in enumerate(arquivos_pdf):
        data_str = arquivo[:8]
        if data_str.isdigit():
            data_formatada = f"{data_str[6:8]}/{data_str[4:6]}/{data_str[:4]}"
        else:
            data_formatada = "Data Inválida"

        try:
            contadas, a_verificar, detalhes_do_dia = analisar_documento(os.path.join(pasta_diarios, arquivo), busca)
            paginas_do_dia = contadas | a_verificar

            if paginas_do_dia:
                if data_formatada == "Data Inválida":
                    erros.append(f"{arquivo}: o nome não começa com a data no formato AAAAMMDD; a Aba 2 não conseguirá separar suas páginas.")
                resultados.append({
                    "DATA": data_formatada,
                    "PORTARIA": "CAI, DAI, DAS, APOSTILAS",
                    "PÁGINAS": ", ".join(str(p) for p in sorted(paginas_do_dia)),
                    "PÁGINAS A VERIFICAR": ", ".join(str(p) for p in sorted(a_verificar)),
                    "FOLHA": "PROCESSADO"
                })
            for d in detalhes_do_dia:
                detalhes.append({
                    "DATA": data_formatada,
                    "PÁGINA": ", ".join(str(p) for p in d["paginas"]),
                    "TIPO": d["tipo"],
                    "RESULTADO": d["resultado"],
                    "MOTIVO": d["motivo"],
                    "REFERÊNCIA": d["referencia"],
                    "TRECHO": d["trecho"],
                    "ARQUIVO": arquivo
                })
        except Exception as e:
            erros.append(f"Erro ao processar {arquivo}: {e}")

        barra_progresso.progress((i + 1) / len(arquivos_pdf))

    if not resultados:
        return None, None, "Nenhuma ocorrência encontrada nos PDFs.", erros

    df = pd.DataFrame(resultados)
    df_detalhes = pd.DataFrame(detalhes)
    # A primeira aba continua sendo o índice lido pelas Abas 2 e 3
    with pd.ExcelWriter(arquivo_saida) as writer:
        df.to_excel(writer, sheet_name="Índice", index=False)
        df_detalhes.to_excel(writer, sheet_name="Detalhes", index=False)
    estilizar_planilha(arquivo_saida)
    return df, df_detalhes, "Sucesso", erros

def separar_pdfs(caminho_planilha, diretorio_integra, diretorio_destino, sobrescrever=False):
    df, erro = ler_planilha_indice(caminho_planilha)
    if erro:
        return False, [erro]

    if not os.path.exists(diretorio_integra):
        return False, [f"Diretório de PDFs na íntegra não encontrado: {diretorio_integra}"]

    pdfs_integra = listar_pdfs(diretorio_integra)
    logs = []

    for _, row in df.iterrows():
        lista_paginas = extrair_lista_paginas(row['PÁGINAS'])
        if not lista_paginas:
            continue

        data = converter_data(row['DATA'])
        if not data:
            logs.append(f"Aviso: linha com data inválida ignorada ({row['DATA']}).")
            continue

        nome_pasta_dia = f"{data.day:02d} {MESES_EXTENSO[data.month]}"
        prefixo_data = data.strftime("%Y%m%d")

        pdf_principal = next((f for f in pdfs_integra if f.startswith(prefixo_data)), None)
        if not pdf_principal:
            logs.append(f"Aviso: PDF na íntegra com prefixo {prefixo_data} não foi encontrado em {diretorio_integra}.")
            continue

        caminho_pasta_dia = os.path.join(diretorio_destino, nome_pasta_dia)
        os.makedirs(caminho_pasta_dia, exist_ok=True)

        try:
            with fitz.open(os.path.join(diretorio_integra, pdf_principal)) as doc:
                for num_pagina in lista_paginas:
                    idx_pagina = num_pagina - 1
                    if not 0 <= idx_pagina < len(doc):
                        logs.append(f"Erro: Página {num_pagina} está fora dos limites do documento {pdf_principal}.")
                        continue

                    nome_saida_pdf = f"{num_pagina:02d}.pdf"
                    caminho_saida = os.path.join(caminho_pasta_dia, nome_saida_pdf)
                    if os.path.exists(caminho_saida) and not sobrescrever:
                        logs.append(f"Mantida (já existia): {nome_pasta_dia} -> {nome_saida_pdf}")
                        continue

                    with fitz.open() as novo_doc:
                        novo_doc.insert_pdf(doc, from_page=idx_pagina, to_page=idx_pagina)
                        novo_doc.save(caminho_saida)
                    logs.append(f"Extraída com sucesso: {nome_pasta_dia} -> {nome_saida_pdf}")
        except Exception as e:
            logs.append(f"Erro ao processar o arquivo {pdf_principal}: {e}")

    return True, logs

def distribuir_paginas(caminho_planilha, nomes_equipe):
    if not nomes_equipe:
        return None, "Informe ao menos um membro na Equipe de Revisão (menu lateral)."

    df, erro = ler_planilha_indice(caminho_planilha)
    if erro:
        return None, erro

    todas_paginas = []
    for _, row in df.iterrows():
        data = formatar_data(row['DATA'])
        for p in extrair_lista_paginas(row['PÁGINAS']):
            todas_paginas.append((data, p))

    total_paginas = len(todas_paginas)
    if total_paginas == 0:
        return None, "Nenhuma página para distribuir na planilha informada."

    # Divisão equilibrada: as páginas que sobram vão uma para cada membro, a partir do primeiro,
    # de modo que ninguém fique com mais de uma página de diferença em relação aos demais.
    base_qtd, sobra = divmod(total_paginas, len(nomes_equipe))
    distribuicao = {}
    inicio = 0
    for i, nome in enumerate(nomes_equipe):
        qtd = base_qtd + (1 if i < sobra else 0)
        distribuicao[nome] = todas_paginas[inicio : inicio + qtd]
        inicio += qtd

    texto_saida = ""
    for nome in nomes_equipe:
        qtd = len(distribuicao[nome])
        texto_saida += f"**{nome}** ({qtd} {'página' if qtd == 1 else 'páginas'})\n"
        paginas_por_data = {}
        for data, pag in distribuicao[nome]:
            paginas_por_data.setdefault(data, []).append(str(pag))

        for data, pags in paginas_por_data.items():
            texto_saida += f"- {data}: {', '.join(pags)}\n"
        texto_saida += "\n"

    return texto_saida, total_paginas

def extrair_texto_portarias(pasta_diarios, arquivo_saida_txt):
    """Retorna (quantidade de portarias, mensagem, lista de erros por arquivo)."""
    if not os.path.exists(pasta_diarios):
        return 0, f"A pasta '{pasta_diarios}' não existe.", []

    arquivos_pdf = listar_pdfs(pasta_diarios)
    if not arquivos_pdf:
        return 0, "Nenhum PDF na íntegra encontrado.", []

    contador_total = 0
    erros = []
    barra_progresso = st.progress(0)

    with open(arquivo_saida_txt, 'w', encoding='utf-8') as arquivo_txt:
        for i, arquivo in enumerate(arquivos_pdf):
            try:
                atos, _ = dividir_em_atos(ler_paginas_pdf(os.path.join(pasta_diarios, arquivo)))
                for ato in atos:
                    if ato.tipo != "portaria":
                        continue
                    bloco = ato.texto
                    if TERMOS_CARGOS.search(bloco) and TERMOS_ACAO.search(bloco) and TERMO_MATRICULA.search(bloco):
                        bloco_limpo = re.sub(r'\s+', ' ', bloco).strip()
                        arquivo_txt.write(bloco_limpo + "\n\n\n")
                        contador_total += 1
            except Exception as e:
                erros.append(f"Erro ao processar {arquivo}: {e}")

            barra_progresso.progress((i + 1) / len(arquivos_pdf))

    return contador_total, "Sucesso", erros

# ==========================================
# INTERFACE DO USUÁRIO (UI)
# ==========================================

st.title("📄 Sistema de Processamento - DJE")
st.markdown("Automação de extração de portarias de requisitados, separação de páginas e distribuição de metas da equipe.")

# A Aba 1 grava aqui a planilha que acabou de gerar; ela é aplicada antes de o seletor ser desenhado
if 'planilha_recem_gerada' in st.session_state:
    st.session_state.planilha_ativa = st.session_state.pop('planilha_recem_gerada')

# --- MENU LATERAL (CONFIGURAÇÕES) ---
with st.sidebar:
    st.header("⚙️ Configurações de Diretórios")
    diretorio_integra = resolver_caminho(st.text_input("Pasta dos PDFs na Íntegra", value="./DO"))
    diretorio_separados = resolver_caminho(st.text_input("Pasta dos PDFs Separados", value="./DJE_Separados"))
    diretorio_acervo = resolver_caminho(st.text_input(
        "Pasta do Acervo (DJEs anteriores)",
        value="./Acervo",
        help="DJEs de datas anteriores, com o mesmo padrão de nome (AAAAMMDD...). Usado para consultar a "
             "portaria original citada em apostilas quando ela não está na pasta dos PDFs na Íntegra."
    ))

    st.divider()
    st.header("🗂️ Padrão de Nomenclatura")
    planilha_base = st.text_input("Nome Base da Planilha", value="DO-exp_Processado_Agrupado.xlsx")
    txt_base = st.text_input("Nome Base do TXT", value="Portarias_requisitados.txt")

    # Seletor da planilha usada nas Abas 2 e 3: lista os arquivos da pasta, então sobrevive a recarregar a página
    # e permite usar uma planilha antiga ou corrigida à mão no Excel.
    planilhas_disponiveis = listar_planilhas()
    if planilhas_disponiveis:
        if st.session_state.get('planilha_ativa') not in planilhas_disponiveis:
            st.session_state.planilha_ativa = planilhas_disponiveis[0]
        st.selectbox(
            "📌 Planilha usada nas Abas 2 e 3",
            planilhas_disponiveis,
            key='planilha_ativa',
            help="Lista as planilhas .xlsx da pasta do sistema, da mais recente para a mais antiga."
        )
        caminho_planilha_ativa = os.path.join(BASE_DIR, st.session_state.planilha_ativa)
    else:
        st.info("📌 Nenhuma planilha de índice encontrada. Gere uma na Aba 1.")
        caminho_planilha_ativa = None

    st.divider()
    st.header("👥 Equipe de Revisão")
    equipe_input = st.text_area("Membros (um por linha)", value="Álvaro\nSidnei\nCatiane\nFred")
    nomes_equipe = [nome.strip() for nome in equipe_input.split('\n') if nome.strip()]

nome_planilha_ativa = os.path.basename(caminho_planilha_ativa) if caminho_planilha_ativa else "nenhuma"

# --- ABAS DE NAVEGAÇÃO ---
tab1, tab2, tab3, tab4 = st.tabs([
    "📊 1. Indexação (Excel)",
    "✂️ 2. Separação de Páginas",
    "👥 3. Distribuição",
    "📝 4. Extração de Texto (Requisitados)"
])

with tab1:
    st.subheader("Gerar Planilha de Ocorrências no DJE")
    st.write(
        "Varre os arquivos na íntegra buscando portarias com incidência de cargos (CAI, DAI, DAS) e ações administrativas, "
        "e apostilas (ou atos que tornam portarias sem efeito) com efeito financeiro. Quando a apostila não é clara, "
        "a portaria original é consultada na pasta dos PDFs na Íntegra e no Acervo."
    )
    if not os.path.isdir(diretorio_acervo):
        st.caption(f"ℹ️ A pasta do Acervo (`{diretorio_acervo}`) não existe: apostilas que citam DJEs fora do período ficarão como \"a verificar\".")

    if st.button("▶️ Mapear e Gerar Excel", type="primary"):
        with st.spinner("Analisando páginas do DJE..."):
            arquivo_saida = nome_com_timestamp(planilha_base)
            df, df_detalhes, msg, erros = gerar_indice_excel(diretorio_integra, diretorio_acervo, os.path.join(BASE_DIR, arquivo_saida))
        # O resultado fica guardado na sessão e a página é recarregada para o seletor do menu lateral
        # já aparecer com a planilha nova selecionada.
        st.session_state.resultado_indice = {"arquivo": arquivo_saida, "df": df, "detalhes": df_detalhes, "msg": msg, "erros": erros}
        if df is not None:
            st.session_state.planilha_recem_gerada = arquivo_saida
        st.rerun()

    resultado = st.session_state.get('resultado_indice')
    if resultado:
        if resultado["df"] is not None:
            st.success(f"Arquivo '{resultado['arquivo']}' gerado com sucesso! As próximas abas já estão configuradas para utilizá-lo.")
            st.dataframe(resultado["df"], width="stretch")

            detalhes = resultado["detalhes"]
            apostilas = detalhes[detalhes["TIPO"] != "Portaria"]
            a_verificar = (resultado["df"]["PÁGINAS A VERIFICAR"] != "").sum()
            if a_verificar:
                st.warning(f"⚠️ {a_verificar} dia(s) têm páginas incluídas por apostilas que precisam de conferência manual "
                           "(coluna \"PÁGINAS A VERIFICAR\"). O motivo está na aba \"Detalhes\" da planilha.")
            with st.expander(f"📎 Apostilas e \"tornar sem efeito\" analisados ({len(apostilas)})"):
                st.dataframe(apostilas, width="stretch", hide_index=True)
        else:
            st.error(resultado["msg"])
        if resultado["erros"]:
            with st.expander(f"⚠️ {len(resultado['erros'])} aviso(s) durante a leitura dos PDFs", expanded=True):
                for erro in resultado["erros"]:
                    st.text(erro)

with tab2:
    st.subheader("Separação de Páginas Específicas")
    st.write(f"Extrai as páginas identificadas no índice atual (`{nome_planilha_ativa}`) e organiza na pasta de destino.")
    sobrescrever = st.checkbox(
        "Sobrescrever páginas já separadas",
        value=False,
        help="Desmarcado, as páginas que já existem na pasta de destino são mantidas como estão."
    )

    if st.button("▶️ Separar PDFs"):
        with st.spinner("Extraindo e organizando páginas em novos arquivos..."):
            sucesso, logs = separar_pdfs(caminho_planilha_ativa, diretorio_integra, diretorio_separados, sobrescrever)
            if sucesso:
                st.success(f"Processo de separação finalizado! Os arquivos foram salvos em '{diretorio_separados}'.")
                with st.expander("Ver Detalhes do Log de Separação"):
                    for log in logs:
                        st.text(log)
            else:
                st.error("\n".join(logs))

with tab3:
    st.subheader("Distribuição da Carga de Trabalho")
    st.write(f"Divide de forma equilibrada as páginas mapeadas no índice atual (`{nome_planilha_ativa}`) entre a equipe.")

    if st.button("▶️ Calcular Distribuição"):
        texto_dist, total = distribuir_paginas(caminho_planilha_ativa, nomes_equipe)
        if texto_dist:
            st.success(f"Divisão concluída! {total} páginas distribuídas entre {len(nomes_equipe)} servidores.")
            st.markdown(texto_dist)
        else:
            st.error(total)

with tab4:
    st.subheader("Consolidação de Portarias de Requisitados")
    st.write("Filtra e extrai o texto literal exclusivamente das portarias associadas a servidores requisitados (matrículas iniciadas com 5).")

    if st.button("▶️ Extrair Portarias (TXT)"):
        with st.spinner("Varrendo documentos e isolando atos de requisitados..."):
            arquivo_txt_saida = nome_com_timestamp(txt_base)
            caminho_txt_saida = os.path.join(BASE_DIR, arquivo_txt_saida)

            total_portarias, msg, erros = extrair_texto_portarias(diretorio_integra, caminho_txt_saida)
            if total_portarias > 0:
                st.success(f"Processo concluído! {total_portarias} portarias foram compiladas no arquivo '{arquivo_txt_saida}'.")

                with open(caminho_txt_saida, "rb") as f:
                    st.download_button(
                        label="📥 Baixar Portarias Consolidadas (TXT)",
                        data=f.read(),
                        file_name=arquivo_txt_saida,
                        mime="text/plain"
                    )
            else:
                st.warning(f"Nenhuma portaria de requisitados foi extraída. Detalhes: {msg}")

            if erros:
                with st.expander(f"⚠️ {len(erros)} PDF(s) não puderam ser lidos", expanded=True):
                    for erro in erros:
                        st.text(erro)
